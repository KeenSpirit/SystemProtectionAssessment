"""
Legacy SEQ fault level workbook layout and save.

Formerly legacy_script/save_results.py; renamed so it cannot be
confused with the save_results package. Layout unchanged.

Functions:
    output_results: Build the workbook from the bridged results.
    save_results: Save it to the protection drive, falling back to the
        user's home folder.
"""


def output_results(app, sub_name, external_grid, feeders_devices_inrush, results_max_3p, results_max_2p,
            results_max_pg, results_min_2p, results_min_3p, results_min_pg, result_sys_norm_min_2p,
            result_sys_norm_min_pg, feeders_sections_trmax_size, results_max_tr_3p, results_max_tr_pg,
            results_all_max_3p, results_all_max_2p, results_all_max_pg, results_all_min_3p, results_all_min_2p,
            results_all_min_pg, result_all_sys_norm_min_2p, result_all_sys_norm_min_pg, feeders_devices_load,
            results_lines_max_3p, results_lines_max_2p, results_lines_max_pg, results_lines_min_3p,
            results_lines_min_2p, results_lines_min_pg, result_lines_sys_norm_min_2p, result_lines_sys_norm_min_pg,
            result_lines_type, result_lines_therm_rating, fdrs_open_switches):
    """Format results file for MS Excel"""

    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    import datetime

    app.PrintPlain("Creating output file...")

    wb = openpyxl.Workbook()

    # Set up Inputs sheet
    wb.create_sheet(index=0, title="Inputs Summary")
    wb.remove(wb["Sheet"])
    i_s = wb["Inputs Summary"]
    i_s.column_dimensions['A'].width = 19.86
    i_s.column_dimensions['B'].width = 9.14
    i_s.column_dimensions['C'].width = 9.14

    # set fonts
    heading_font = Font(size=14, bold=True)
    subheading_font = Font(size=11, bold=True)
    i_s['A1'].font = heading_font
    i_s['A3'].font = subheading_font
    i_s['A7'].font = heading_font
    i_s['A13'].font = subheading_font
    i_s['A22'].font = subheading_font

    # Set headings
    i_s['A1'] = sub_name + " Feeder Fault Level Study"
    i_s['A3'] = "Script Run Date:"
    i_s['A4'] = datetime.datetime.now()
    i_s['A7'] = "Input Summary"
    i_s['A9'] = "Short-circuit calculation method: Complete"
    i_s['A10'] = "c-Factor (max): 1.1"
    i_s['C10'] = "Voltage factor c (max): 1.1"
    i_s['A11'] = "c-Factor (min): 1.0"
    i_s['C11'] = "Voltage factor c (min): 1.0"
    i_s['A13'] = "External Grid Data:"
    i_s['A16'] = "3-P fault level (A):"
    i_s['A17'] = "R/X:"
    i_s['A18'] = "Z2/Z1:"
    i_s['A19'] = "X0/X1:"
    i_s['A20'] = "R0/X0:"
    i_s['A22'] = ("NOTE: Open points to adjacent bulk supply substations may not be "
                   "detected unless the relevant grids are active under the current "
                  "project")

    # Write external grid data
    next_column = 2
    for grid, value in external_grid.items():
        i_s.cell(column=next_column, row=14, value=grid.loc_name).font = subheading_font
        i_s.cell(column=next_column, row=15, value="Maximum")
        i_s.cell(column=next_column + 1, row=15, value="Minimum")
        i_s.cell(column=next_column + 2, row=15, value="System Normal Minimum")
        next_row = 16
        for values in value[:5]:
            i_s.cell(column=next_column, row=next_row, value=values)
            next_row += 1
        next_row = 16
        for values in value[5:10]:
            i_s.cell(column=next_column + 1, row=next_row, value=values)
            next_row += 1
        next_row = 16
        for values in value[10:]:
            i_s.cell(column=next_column + 2, row=next_row, value=values)
            next_row += 1
        next_column += 4

    # Set up Results sheet
    wb.create_sheet(index=1, title="Results Summary")
    r_s = wb["Results Summary"]
    length = len(feeders_devices_inrush)
    j = 1
    for i in range(length):
        k = get_column_letter(j)
        r_s.column_dimensions[k].width = 37
        j += 5

    # Set fonts
    r_s.column_dimensions['A'].width = 36.86
    r_s['A1'].font = heading_font
    r_s['A3'].font = heading_font
    r_s['A5'].font = heading_font

    # Set headings
    r_s['A1'] = sub_name + " Feeder Fault Level Study"
    r_s['A3'] = "Results Summary"

    # Set up Detailed Results sheet
    # Set headings and fonts
    for key, value in feeders_devices_inrush.items():
        wb.create_sheet(title=key)
        feed = wb[key]
        feed['A1'] = sub_name + " Feeder Fault Level Study - Detailed Results for " + key
        feed['A1'].font = heading_font

    def nested_dic(a, b, c):
        """Template for writing nested dictionary data to the Results Summary sheet """

        next_column = 2
        for key, value in a.items():
            r_s.cell(column=next_column - 1, row=5, value="FEEDER:").font = heading_font
            r_s.cell(column=next_column, row=5, value=key).font = heading_font
            next_row = 7
            for key2, values2 in value.items():
                r_s.cell(column=next_column - 1, row=next_row, value="SECTION:").font = heading_font
                r_s.cell(column=next_column, row=next_row, value=key2).font = heading_font
                r_s.cell(column=next_column - 1, row=next_row + b, value=c)
                r_s.cell(column=next_column, row=next_row + b, value=values2)
                next_row += 14
            next_column += 5

    def two_nested_dic(a, b, c):
        """Template for writing level 2 nested dictionary data to the Results Summary sheet """

        next_column = 2
        for key, value in a.items():
            next_row = 7
            for key2, value2 in value.items():
                for key3, value3, in value2.items():
                    if key3[-5:] == "_Term":
                        key4 = key3[:-5]
                    else:
                        key4 = key3
                    r_s.cell(column=next_column + 1, row=next_row + b, value=key4)
                    r_s.cell(column=next_column - 1, row=next_row + b, value=c)
                    r_s.cell(column=next_column, row=next_row + b, value=value3)
                next_row += 14
            next_column += 5
        return

    def elements_open(fdrs_open_switches, feeders_devices_inrush):
        """
        """

        feeder_row_dic = {feeder: len(feeders_devices_inrush[feeder]) for feeder in feeders_devices_inrush}

        next_column = 2
        for feeder, open_switches in fdrs_open_switches.items():
            last_row = 7 + (feeder_row_dic[feeder] * 14)
            r_s.cell(column=next_column-1, row=last_row, value="FEEDER OPEN POINTS:").font = heading_font
            if len(open_switches) > 0:
                for site, switch in open_switches.items():
                    switch_name = switch.GetAttribute("loc_name")
                    if switch.GetClassName() == "StaSwitch":
                        r_s.cell(column=next_column, row=last_row, value=switch_name)
                    else:
                        site_name = site.GetAttribute("loc_name")
                        r_s.cell(column=next_column, row=last_row, value=f"{site_name} / {switch_name}")
                    last_row += 1
            else:
                r_s.cell(column=next_column - 1, row=last_row, value="(None detected)")
                last_row += 1
            last_row += 1
            next_column += 5

    def sheets_nested_dic(wb, node_data, line_data, col_node_data, col_line_data, col_name_1, col_name_2):
        """Template for writing level 2 nested dictionary data over multiple sheets
        """

        for key, value in node_data.items():
            feed = wb[key]
            feed.column_dimensions['B'].width = 18.14
            feed.column_dimensions['M'].width = 18.14
            feed.column_dimensions['W'].width = 18.14
            feed.column_dimensions['AG'].width = 18.14
            feed.column_dimensions['AQ'].width = 18.14
            next_column = 3
            for key2, value2 in value.items():
                feed.cell(column=next_column, row=3, value=key2).font = heading_font
                feed.cell(column=next_column - 2, row=3, value="SECTION:").font = heading_font
                feed.cell(column=next_column - 1, row=4, value="Downstream terminal")
                feed.cell(column=next_column + col_node_data, row=4, value=col_name_1)
                next_row = 3
                for key3, value3 in value2.items():
                    if key3[-5:] == "_Term":
                        key4 = key3[:-5]
                    else:
                        key4 = key3
                    feed.cell(column=next_column - 1, row=next_row + 2, value=key4)
                    feed.cell(column=next_column + col_node_data, row=next_row + 2, value=value3)
                    next_row += 1
                feed.cell(column=next_column - 2, row=next_row + 3, value="LINES").font = heading_font
                feed.cell(column=next_column - 1, row=next_row + 4, value="Line")
                feed.cell(column=next_column + col_line_data, row=next_row + 4, value=col_name_2)
                length = 0
                line_dict = line_data[key][key2]
                for key5, value5 in line_dict.items():
                    feed.cell(column=next_column - 1, row=next_row + length + 5, value=key5.loc_name)
                    feed.cell(column=next_column + col_line_data, row=next_row + length + 5, value=value5)
                    length += 1
                next_column += 11
        return

    # Write results to the Results sheet and the Detailed Results sheets in MS Excel
    nested_dic(feeders_devices_inrush, 1, "Inrush (A):")

    two_nested_dic(results_max_3p, 2, "Max 3-P fault level (A) (Site):")
    two_nested_dic(results_max_2p, 3, "Max 2-P fault level (A) (Site):")
    two_nested_dic(results_max_pg, 4, "Max P-G fault level (A) (Site):")
    two_nested_dic(results_min_3p, 5, "Min 3-P fault level (A) (Site):")
    two_nested_dic(results_min_2p, 6, "Min 2-P fault level (A) (Site):")
    two_nested_dic(results_min_pg, 7, "Min P-G fault level (A) (Site):")
    two_nested_dic(result_sys_norm_min_2p, 8, "System normal Min 2-P fault level (A) (Site):")
    two_nested_dic(result_sys_norm_min_pg, 9, "System normal Min P-G fault level (A) (Site):")
    two_nested_dic(feeders_sections_trmax_size, 10, "Largest transformer size (kVA) (Site):")
    two_nested_dic(results_max_tr_3p, 11, "Max 3-P at largest transformer (A) (Site):")
    two_nested_dic(results_max_tr_pg, 12, "Max P-G at largest transformer (A) (Site):")

    elements_open(fdrs_open_switches, feeders_devices_inrush)

    sheets_nested_dic(wb, feeders_devices_load, result_lines_type, -2, -2, "Tfmr load (kVA)", "Conductor type")
    sheets_nested_dic(wb, results_all_max_3p, result_lines_therm_rating, 0, 0, "Max 3P fault", "Rated 1s current (kA)")
    sheets_nested_dic(wb, results_all_max_2p, results_lines_max_3p, 1, 1, "Max 2P fault", "Max 3P fault")
    sheets_nested_dic(wb, results_all_max_pg, results_lines_max_2p, 2, 2, "Max PG fault", "Max 2P fault")
    sheets_nested_dic(wb, results_all_min_3p, results_lines_max_pg, 3, 3, "Min 3P fault", "Max PG fault")
    sheets_nested_dic(wb, results_all_min_2p, results_lines_min_3p, 4, 4, "Min 2P fault", "Min 3P fault")
    sheets_nested_dic(wb, results_all_min_pg, results_lines_min_2p, 5, 5, "Min PG fault", "Min 2P fault")
    sheets_nested_dic(wb, result_all_sys_norm_min_2p, results_lines_min_pg, 6, 6, "Min Sys Norm 2P fault",
                      "Min PG fault")
    sheets_nested_dic(wb, result_all_sys_norm_min_pg, result_lines_sys_norm_min_2p, 7, 7, "Min Sys Norm PG fault",
                      "Min Sys Norm 2P fault")
    sheets_nested_dic(wb, results_all_max_3p, result_lines_sys_norm_min_pg, 0, 8, "Max 3P fault",
                      "Min Sys Norm PG fault")

    return wb


def save_results(app, sub_name, wb):
    """Save file to disk"""

    import os
    from pathlib import Path
    import time

    date_string = time.strftime("%Y%m%d-%H%M%S")
    file_name = sub_name + " fault level study " + date_string + ".xlsx"

    home_path = str(Path.home())
    save_path = r'\\client\Y$\PROTECTION\PowerFactory\Fault Level Studies'
    save_path_2 = '\\\\client\\' + home_path[0] + '$' + home_path[2:]
    save_path_3 = home_path

    # Try to save the output file to the SEQ Protection department drive
    try:
        wb.save(os.path.join(save_path, file_name))
        app.PrintPlain("Output file saved to " + save_path)
    # If this fails, save the output file to the user's home directory
    except Exception:
        try:
            wb.save(os.path.join(save_path_2, file_name))
            app.PrintPlain("Output file saved to " + save_path_2)
        except Exception:
            # When running a local install of PowerFactory
            wb.save(os.path.join(save_path_3, file_name))
            app.PrintPlain("Output file saved to " + save_path_3)

    return
