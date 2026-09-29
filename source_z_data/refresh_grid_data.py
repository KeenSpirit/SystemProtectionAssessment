#!/usr/bin/env python3
"""
refresh_grid_data.py
====================

Rebuild ``grid_results_ee.xlsx`` from a new issue of the Ergon fault level
report, reusing the grid -> reporting bus mapping already held in the
previous ``grid_results_ee.xlsx`` (column C, "Buses").

The expensive part of building the workbook is deciding which reporting
bus(es) each PowerFactory external grid corresponds to. That mapping was
done once by ``build_grid_data.py`` (sub-code, voltage, T/BUS number and
fault level matching, plus manual review) and is stored in column C. This
script keeps that mapping and only refreshes the values:

    Bound = Max    <- 'Max-Max Fault Level Report' tab, highest 3P fault
                      of the mapped buses
    Bound = Min    <- 'Min Fault Level Report' tab, lowest 3P fault of
                      the mapped buses
    Bound = Min_SN <- duplicate of the Min row

Columns F..J (3P fault kA, R/X, Z2/Z1, X0/X1, R0/X0) come from report
columns V..Z, exactly as ``build_grid_data.value_row`` reads them.

Grids with no mapped bus are NOT written to 'Grid Results'. Their rows in
earlier workbooks were copied from ``pf_external_grids.xlsx``, i.e. from
the PowerFactory models themselves, in amps. Leaving them out means
``start.get_grid_data`` keeps each model's own value, which is what those
rows held anyway, without the 1000x unit error. They are listed in
'Mapping Notes' so they can be mapped later through --overrides.

If a mapped bus no longer appears in the new report (renamed or removed),
the grid is re-matched with ``build_grid_data.match_grid``, using its
previous Max 3P fault to choose between candidate buses. The outcome is
flagged in 'Mapping Notes' either way.

Usage
-----
    python -m source_z_data.refresh_grid_data \
        --report    "2027 Fault Level Report (Ergon - Internal)_V1_0.xlsx" \
        --previous  data/grid_results_ee.xlsx \
        --out       data/grid_results_ee_2027.xlsx \
        [--overrides grid_bus_overrides.csv]

--overrides is an optional CSV with columns Grid,Buses (Buses separated by
", "), used to add a mapping for an unmapped grid or replace an existing
one. An empty Buses value removes a grid's mapping.

The script only reads and writes Excel files; it needs no PowerFactory.
"""

from __future__ import annotations

import argparse
import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font

try:
    from source_z_data.build_grid_data import (
        MAX_MATCH_TAB, MAX_VALUE_TAB, MIN_VALUE_TAB, load_report,
        match_grid, value_row,
    )
except ImportError:  # run directly from inside source_z_data/
    from build_grid_data import (  # type: ignore
        MAX_MATCH_TAB, MAX_VALUE_TAB, MIN_VALUE_TAB, load_report,
        match_grid, value_row,
    )

GRID_SHEET = "Grid Results"
NOTES_SHEET = "Mapping Notes"
LOG_SHEET = "Refresh Log"

# Header row, identical to the existing workbook (leading spaces and all),
# so load_source_z_data reads the output unchanged.
GRID_HEADERS = [
    "Bulk Supply Point", " Grid", "Buses", " Bound", "Scenario",
    " 3P fault", " R/X", " Z2/Z1", " X0/X1", " R0/X1",
]
BOUND_ROWS = [
    (" Max", "Maximum", "max"),
    (" Min", "System Normal", "min"),
    (" Min_SN", "System Normal", "min"),
]

# ElmXnet.ikss is in kA; the Ergon report publishes 0.1-40 kA. Anything
# above this is not a kA figure and is never written.
IKSS_KA_LIMIT = 1000.0
# A Max 3P fault change larger than this between issues is flagged.
CHANGE_FLAG = 0.20


@dataclass
class Grid:
    bsp: str                      # raw, as in the previous workbook
    name: str                     # raw, as in the previous workbook
    buses: List[str]
    old_max: Optional[float] = None
    old_rx: Optional[float] = None
    old_min: Optional[float] = None
    old_note: str = ""
    status: str = ""
    notes: List[str] = field(default_factory=list)

    @property
    def key(self) -> str:
        return self.name.strip()


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #

def _split_buses(text) -> List[str]:
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return []
    return [b.strip() for b in str(text).split(",") if b.strip()]


def _num(value) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_previous(path: str | Path) -> Dict[str, Grid]:
    """Grids of the previous workbook, in their original order."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    grids: Dict[str, Grid] = {}
    try:
        for row in wb[GRID_SHEET].iter_rows(min_row=2, max_col=10, values_only=True):
            if not row or row[1] is None:
                continue
            bsp, name, buses, bound = row[0], str(row[1]), row[2], str(row[3] or "").strip()
            g = grids.setdefault(name.strip(), Grid(str(bsp or ""), name, _split_buses(buses)))
            f3p = _num(row[5])
            # Only report-sourced rows (buses present) are known to be in kA.
            if g.buses and bound == "Max":
                g.old_max, g.old_rx = f3p, _num(row[6])
            elif g.buses and bound == "Min":
                g.old_min = f3p

        if NOTES_SHEET in wb.sheetnames:
            for row in wb[NOTES_SHEET].iter_rows(min_row=2, max_col=7, values_only=True):
                if row and row[1] is not None and str(row[1]).strip() in grids:
                    grids[str(row[1]).strip()].old_note = str(row[6] or "")
    finally:
        wb.close()
    return grids


def read_overrides(path: Optional[str | Path]) -> Dict[str, List[str]]:
    if not path:
        return {}
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return {
            r["Grid"].strip(): _split_buses(r.get("Buses"))
            for r in csv.DictReader(fh)
            if r.get("Grid", "").strip()
        }


# --------------------------------------------------------------------------- #
# Refresh
# --------------------------------------------------------------------------- #

def refresh(report_path, previous_path, out_path, overrides_path=None) -> dict:
    grids = read_previous(previous_path)
    overrides = read_overrides(overrides_path)
    mx = load_report(report_path, MAX_MATCH_TAB)
    mm = load_report(report_path, MAX_VALUE_TAB)
    mn = load_report(report_path, MIN_VALUE_TAB)
    known = set(mm["D"]) & set(mn["D"])

    unknown_overrides = sorted(set(overrides) - set(grids))
    rows, log = [], []
    counts = {"written": 0, "unchanged_mapping": 0, "override": 0,
              "rematched": 0, "not_written": 0, "flagged_change": 0}

    for g in grids.values():
        if g.key in overrides:
            g.buses = overrides[g.key]
            g.status = "override"
            g.notes.append("mapping from overrides file")
            counts["override"] += 1

        if not g.buses:
            g.status = g.status or "not written"
            g.notes.append(
                "no reporting bus mapped; not written, so the model keeps its "
                "own grid values" + (f" (previous note: {g.old_note})" if g.old_note else "")
            )
            counts["not_written"] += 1
            continue

        missing = [b for b in g.buses if b not in known]
        if missing:
            present = [b for b in g.buses if b in known]
            if present:
                g.notes.append(f"bus(es) no longer in report, dropped: {', '.join(missing)}")
                g.buses = present
            else:
                f3p_amps = None if g.old_max is None else g.old_max * 1000.0
                m = match_grid(g.key, mx, f3p_amps, g.old_rx)
                if m["status"] == "matched":
                    g.notes.append(
                        f"re-matched ({m['confidence']}): previous bus(es) "
                        f"{', '.join(missing)} not in report; {m['note']}"
                    )
                    g.buses = m["buses"]
                    g.status = "rematched"
                    counts["rematched"] += 1
                else:
                    g.notes.append(
                        f"previous bus(es) {', '.join(missing)} not in report and "
                        f"re-match {m['status']}: {m['note']}; not written"
                    )
                    g.status = "not written"
                    counts["not_written"] += 1
                    continue

        vals_max, w1 = value_row(mm, g.buses, "max")
        vals_min, w2 = value_row(mn, g.buses, "min")
        g.notes += [w for w in (w1, w2) if w]

        bad = [
            v for v in (vals_max[0], vals_min[0])
            if v is None or pd.isna(v) or not (0 < float(v) <= IKSS_KA_LIMIT)
        ]
        if bad:
            g.notes.append(f"3P fault outside 0-{IKSS_KA_LIMIT:g} kA ({bad}); not written")
            g.status = "not written"
            counts["not_written"] += 1
            continue

        for bound, scenario, which in BOUND_ROWS:
            vals = vals_max if which == "max" else vals_min
            rows.append([g.bsp, g.name, ", ".join(g.buses), bound, scenario,
                         *[None if pd.isna(v) else float(v) for v in vals]])
        counts["written"] += 1
        if not g.status:
            g.status = "refreshed"
            counts["unchanged_mapping"] += 1

        change = None
        if g.old_max:
            change = (float(vals_max[0]) - g.old_max) / g.old_max
            if abs(change) > CHANGE_FLAG:
                g.notes.append(f"Max 3P fault changed {change:+.0%} since the previous issue")
                counts["flagged_change"] += 1
        log.append([g.bsp.strip(), g.key, ", ".join(g.buses), g.old_max,
                    float(vals_max[0]), g.old_min, float(vals_min[0]),
                    None if change is None else round(change, 4), g.status])

    notes = [
        [g.bsp.strip(), g.key, ", ".join(g.buses) or None, g.status, "; ".join(g.notes)]
        for g in grids.values()
        if g.status != "refreshed" or g.notes
    ]
    for name in unknown_overrides:
        notes.append([None, name, ", ".join(overrides[name]), "override ignored",
                      "grid not in the previous workbook; add it there first"])

    write_workbook(out_path, rows, notes, log)
    counts["grids"] = len(grids)
    counts["rows"] = len(rows)
    return counts


# --------------------------------------------------------------------------- #
# Writing
# --------------------------------------------------------------------------- #

def _sheet(wb, title, headers, data, widths=None, wrap_col=None):
    ws = wb.create_sheet(title)
    bold, body = Font(name="Arial", size=10, bold=True), Font(name="Arial", size=10)
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h).font = bold
    for r, values in enumerate(data, 2):
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = body
            if wrap_col and c == wrap_col:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(data) + 1}"
    for col, w in (widths or {}).items():
        ws.column_dimensions[col].width = w
    return ws


def write_workbook(out_path, rows, notes, log) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, GRID_SHEET, GRID_HEADERS, rows,
           {"A": 22, "B": 30, "C": 45, "E": 15})
    _sheet(wb, NOTES_SHEET, ["Bulk Supply Point", "Grid", "Mapped Bus(es)", "Status", "Note"],
           notes, {"A": 22, "B": 34, "C": 45, "D": 14, "E": 110}, wrap_col=5)
    _sheet(wb, LOG_SHEET,
           ["Bulk Supply Point", "Grid", "Bus(es)", "Previous Max 3P (kA)", "New Max 3P (kA)",
            "Previous Min 3P (kA)", "New Min 3P (kA)", "Max change", "Status"],
           log, {"A": 22, "B": 30, "C": 45})
    wb.save(out_path)


# --------------------------------------------------------------------------- #

def main() -> None:
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--report", required=True, help="Ergon fault level report (.xlsx)")
    p.add_argument("--previous", required=True, help="existing grid_results_ee.xlsx")
    p.add_argument("--out", required=True, help="output workbook")
    p.add_argument("--overrides", help="optional CSV of Grid,Buses mappings")
    a = p.parse_args()
    if Path(a.out).resolve() == Path(a.previous).resolve():
        p.error("--out must differ from --previous")

    c = refresh(a.report, a.previous, a.out, a.overrides)
    print(
        f"{c['grids']} grids in previous workbook\n"
        f"  written (3 rows each)     : {c['written']}  ({c['rows']} rows)\n"
        f"    mapping reused          : {c['unchanged_mapping']}\n"
        f"    from overrides          : {c['override']}\n"
        f"    re-matched              : {c['rematched']}\n"
        f"  not written (model kept)  : {c['not_written']}\n"
        f"  Max 3P change > {CHANGE_FLAG:.0%}      : {c['flagged_change']}\n"
        f"written to {a.out}"
    )


if __name__ == "__main__":
    main()